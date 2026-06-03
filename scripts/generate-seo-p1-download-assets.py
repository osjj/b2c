from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


OUT_DIR = Path("public/downloads")
OUT_DIR.mkdir(parents=True, exist_ok=True)


def styles():
    base = getSampleStyleSheet()
    base.add(
        ParagraphStyle(
            name="Tiny",
            parent=base["Normal"],
            fontSize=7.5,
            leading=9,
            textColor=colors.HexColor("#4b5563"),
        )
    )
    base.add(
        ParagraphStyle(
            name="SectionTitle",
            parent=base["Heading2"],
            fontSize=13,
            leading=16,
            spaceBefore=12,
            spaceAfter=6,
            textColor=colors.HexColor("#111827"),
        )
    )
    return base


def add_table(story, headers, rows, widths):
    data = [[Paragraph(f"<b>{cell}</b>", styles()["Tiny"]) for cell in headers]]
    data.extend([[Paragraph(str(cell), styles()["Tiny"]) for cell in row] for row in rows])
    table = Table(data, colWidths=widths, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e5e7eb")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#111827")),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    story.append(table)
    story.append(Spacer(1, 10))


def build_rfq_pdf():
    path = OUT_DIR / "construction-ppe-rfq-template.pdf"
    doc = SimpleDocTemplate(
        str(path),
        pagesize=letter,
        rightMargin=0.55 * inch,
        leftMargin=0.55 * inch,
        topMargin=0.55 * inch,
        bottomMargin=0.55 * inch,
    )
    s = styles()
    story = [
        Paragraph("Construction PPE RFQ Template", s["Title"]),
        Paragraph(
            "Use this template when requesting quotes for construction PPE across roles, hazards, standards, sizes, packaging, and delivery requirements.",
            s["Normal"],
        ),
        Spacer(1, 10),
    ]
    add_table(
        story,
        ["Project field", "Buyer input"],
        [
            ["Project / site name", ""],
            ["Delivery country / market", ""],
            ["Target standards", "ANSI / ASTM / NIOSH / EN / CE / buyer-specific"],
            ["Worker roles", "General labor, road crew, demolition, scaffolding, operators, electrical, visitors"],
            ["Expected order quantity", ""],
            ["Quotation deadline", ""],
            ["Sample / document requirements", "Certificates, test reports, packaging photos, samples, labels"],
        ],
        [2.05 * inch, 5.15 * inch],
    )
    story.append(Paragraph("PPE Category Scope", s["SectionTitle"]))
    add_table(
        story,
        ["Category", "Required standard / rating", "Quantity", "Size range", "Notes"],
        [
            ["Head protection", "ANSI Z89.1 Type/Class or EN 397", "", "", ""],
            ["Eye / face protection", "ANSI Z87.1 or EN 166", "", "", ""],
            ["Hand protection", "ANSI/ISEA cut or EN 388", "", "", ""],
            ["Foot protection", "ASTM F2413 or EN ISO 20345", "", "", ""],
            ["High-visibility clothing", "ANSI/ISEA 107 or EN ISO 20471", "", "", ""],
            ["Respiratory protection", "NIOSH approval / cartridge or filter", "", "", ""],
            ["Hearing protection", "NRR / SNR", "", "", ""],
            ["Fall protection", "ANSI / OSHA / EN system requirements", "", "", ""],
        ],
        [1.45 * inch, 1.75 * inch, 0.8 * inch, 1.0 * inch, 2.2 * inch],
    )
    story.append(Paragraph("Supplier Response Checklist", s["SectionTitle"]))
    add_table(
        story,
        ["Supplier item", "Requested response"],
        [
            ["Unit price and MOQ", ""],
            ["Lead time and sample availability", ""],
            ["Certification documents", ""],
            ["Packaging and carton labeling", ""],
            ["Size replenishment support", ""],
            ["Recommended substitutions", ""],
            ["Validity period of quote", ""],
        ],
        [2.05 * inch, 5.15 * inch],
    )
    doc.build(story)


def build_checklist_pdf():
    path = OUT_DIR / "construction-ppe-checklist.pdf"
    doc = SimpleDocTemplate(
        str(path),
        pagesize=letter,
        rightMargin=0.55 * inch,
        leftMargin=0.55 * inch,
        topMargin=0.55 * inch,
        bottomMargin=0.55 * inch,
    )
    s = styles()
    story = [
        Paragraph("Construction PPE Checklist", s["Title"]),
        Paragraph(
            "Field checklist for site induction, supervisor pre-start checks, and procurement planning. Match final PPE to the hazard assessment and local requirements.",
            s["Normal"],
        ),
        Spacer(1, 10),
    ]
    add_table(
        story,
        ["Check", "Yes / No", "Notes"],
        [
            ["Baseline site PPE issued: head, eye, hand, foot, and hi-vis", "", ""],
            ["PPE selected to properly fit each affected worker", "", ""],
            ["Task hazards reviewed: fall, traffic, silica, noise, heat, electrical, splash, cuts", "", ""],
            ["Respirator users covered by program, fit testing, training, and replacement filters", "", ""],
            ["Fall protection inspected and matched to anchor / rescue plan", "", ""],
            ["Safety footwear certified and size range confirmed", "", ""],
            ["Gloves selected by task: cut, impact, coating, chemical, heat, or electrical", "", ""],
            ["Eye and face protection compatible with helmet, respirator, and task", "", ""],
            ["Hi-vis class selected by traffic, vehicle movement, light level, and weather", "", ""],
            ["Replacement stock available for consumables and damaged items", "", ""],
        ],
        [4.0 * inch, 0.9 * inch, 2.3 * inch],
    )
    story.append(PageBreak())
    story.append(Paragraph("Role-Based Add-On Matrix", s["SectionTitle"]))
    add_table(
        story,
        ["Role / task", "Add-on PPE to confirm", "Procurement note"],
        [
            ["Road and bridge crew", "Class 2/Class 3 hi-vis, hearing, eye, footwear, fall protection for bridge edges", "Separate day, night, rain, and heat needs."],
            ["Demolition / concrete cutting", "Respiratory protection, sealed eyewear, face shield, hearing, cut gloves, puncture boots", "Confirm silica controls before respirator line items."],
            ["Scaffolding / height work", "Harness, connector, helmet retention, anti-slip boots, tool lanyards", "Do not mix scaffold erector and scaffold user kits."],
            ["Heavy equipment operators", "Cab kit, hi-vis, hearing, boots, gloves, glasses, dust review", "Plan for outside-cab tasks and walk-around checks."],
            ["Electrical installation", "Class E head protection, insulating gloves, EH footwear, arc-rated layers where required", "Treat electrical PPE as a program item."],
        ],
        [1.55 * inch, 3.25 * inch, 2.4 * inch],
    )
    doc.build(story)


def style_sheet(ws):
    header_fill = PatternFill("solid", fgColor="111827")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for idx, width in enumerate([26, 24, 22, 16, 16, 18, 24, 28, 28], start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width


def build_xlsx():
    path = OUT_DIR / "ppe-size-standards-planning-sheet.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "RFQ Inputs"
    ws.append(
        [
            "Role / crew",
            "PPE category",
            "Product requirement",
            "Standard / rating",
            "Workers",
            "Per worker",
            "Replacement factor",
            "Estimated qty",
            "Notes",
        ]
    )
    rows = [
        ["General site entry", "Head protection", "Safety helmet or hard hat", "ANSI Z89.1 / EN 397", 50, 1, 1.1, "=E2*F2*G2", "Confirm Type I/II and Class G/E/C"],
        ["Road crew", "Hi-vis", "Vest, jacket, or full set", "ANSI/ISEA 107 / EN ISO 20471", 30, 1, 1.2, "=E3*F3*G3", "Separate day/night and rainwear"],
        ["Demolition", "Respiratory", "Respirator and filters", "NIOSH approval", 20, 1, 1.3, "=E4*F4*G4", "Fit testing and filter stock required"],
        ["Scaffolding", "Fall protection", "Harness and connector", "ANSI / EN as required", 12, 1, 1.15, "=E5*F5*G5", "Do not mix with general kit"],
    ]
    for row in rows:
        ws.append(row)
    style_sheet(ws)
    ws["K1"] = "Total estimated qty"
    ws["K1"].font = Font(bold=True)
    ws["K2"] = "=SUM(H2:H200)"
    ws["K2"].number_format = "0"
    ws.column_dimensions["K"].width = 20

    standards = wb.create_sheet("Standards Matrix")
    standards.append(["PPE category", "US reference", "EU / international reference", "Buyer verification"])
    for row in [
        ["Head protection", "ANSI/ISEA Z89.1", "EN 397", "Type, class, venting, chin strap, suspension"],
        ["Eye / face", "ANSI Z87.1", "EN 166", "Lens type, side protection, anti-fog, face shield pairing"],
        ["Footwear", "ASTM F2413", "EN ISO 20345", "Toe, puncture, slip, EH/SD, size curve"],
        ["Hi-vis", "ANSI/ISEA 107", "EN ISO 20471", "Type/class, color, reflective layout, weather layer"],
        ["Respiratory", "NIOSH approval", "EN 149 / EN 140 where applicable", "Filter/cartridge, fit testing, storage, change-out"],
        ["Hearing", "NRR", "EN 352 / SNR", "Earplug/earmuff, helmet compatibility, communication"],
        ["Gloves", "ANSI/ISEA cut or impact", "EN 388", "Cut, impact, coating, chemical, heat, dexterity"],
        ["Fall protection", "ANSI / OSHA requirements", "EN 361 / EN 355 / EN 360", "Harness fit, connector, anchor, rescue plan"],
    ]:
        standards.append(row)
    style_sheet(standards)

    size = wb.create_sheet("Size Plan")
    size.append(["PPE item", "XS", "S", "M", "L", "XL", "2XL", "3XL", "Notes"])
    for row in [
        ["Hi-vis vest / jacket", "", "", "", "", "", "", "", "Split by worker size curve and climate layers"],
        ["Gloves", "", "", "", "", "", "", "", "Separate coated, cut, impact, and chemical glove sizes"],
        ["Footwear", "", "", "", "", "", "", "", "Use actual shoe sizes and width needs, not alpha sizes"],
        ["Harness", "", "", "", "", "", "", "", "Confirm weight range and fit-critical adjustments"],
    ]:
        size.append(row)
    style_sheet(size)

    replacement = wb.create_sheet("Replacement Stock")
    replacement.append(["PPE category", "Normal replacement trigger", "Recommended planning note"])
    for row in [
        ["Eyewear", "Scratched, fogging, broken frame", "Keep point-of-use spare stock"],
        ["Gloves", "Torn, worn coating, contaminated", "Plan by task consumption, not headcount only"],
        ["Respirator filters", "Program change-out schedule or breathing resistance", "Stock by exposure and shift"],
        ["Hard hat suspension", "Damaged, worn, or service-life policy", "Order spare suspensions with shells"],
        ["Fall protection", "Impact event, failed inspection, damage", "Remove from service after fall arrest"],
    ]:
        replacement.append(row)
    style_sheet(replacement)

    wb.save(path)


if __name__ == "__main__":
    build_rfq_pdf()
    build_checklist_pdf()
    build_xlsx()
